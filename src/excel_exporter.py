"""Excel exporter for scraped Facebook data."""

from __future__ import annotations
import io
import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    ProfileData, GroupData, PageData,
    PROFILE_EXCEL_HEADERS, GROUP_EXCEL_HEADERS, PAGE_EXCEL_HEADERS,
)

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="1a73e8", end_color="1a73e8", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Calibri", size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGN = Alignment(vertical="top", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)


def create_excel(profiles: list[ProfileData], groups: list[GroupData], pages: list[PageData]) -> bytes:
    """Create a formatted Excel workbook with separate sheets per entity type."""
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    if profiles:
        _add_sheet(wb, "Profiles", PROFILE_EXCEL_HEADERS, [p.to_excel_row() for p in profiles])
    if groups:
        _add_sheet(wb, "Groups", GROUP_EXCEL_HEADERS, [g.to_excel_row() for g in groups])
    if pages:
        _add_sheet(wb, "Pages", PAGE_EXCEL_HEADERS, [p.to_excel_row() for p in pages])

    if not wb.sheetnames:
        ws = wb.create_sheet("No Data")
        ws["A1"] = "No data was scraped."

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log.info(f"📊 Excel file created ({len(profiles)} profiles, {len(groups)} groups, {len(pages)} pages)")
    return buf.getvalue()


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title)

    # Headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

    # Data rows
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            if isinstance(value, (list, dict)):
                value = str(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = CELL_FONT
            cell.alignment = CELL_ALIGN
            cell.border = THIN_BORDER

    # Auto-fit column widths (approximate)
    for col_idx in range(1, len(headers) + 1):
        max_width = len(headers[col_idx - 1])
        for row_idx in range(2, min(len(rows) + 2, 50)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                max_width = max(max_width, min(len(str(cell.value)), 60))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 4, 65)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
